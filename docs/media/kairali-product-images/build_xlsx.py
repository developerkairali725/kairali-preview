#!/usr/bin/env python3
"""Build kairali_image_catalog.xlsx from mapping.csv."""
import csv, os, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = os.path.dirname(os.path.abspath(__file__))
rows = list(csv.DictReader(open(os.path.join(D, "mapping.csv"), encoding="utf-8")))

COLS = [
    ("new_filename", "proposed_new_filename", 40),
    ("original_drive_filename", "original_drive_filename", 36),
    ("drive_file_id", "drive_file_id", 35),
    ("source_folder", "source_folder", 14),
    ("relative_path", "relative_path", 44),
    ("product_name", "product_name", 34),
    ("site_listed", "site_listed", 11),
    ("evidence_source", "evidence_source", 30),
    ("label_text", "label_text", 50),
    ("matched_source_url", "matched_source_url", 50),
    ("match_confidence", "match_confidence", 17),
    ("description", "description", 58),
    ("tags", "tags", 28),
    ("category", "category", 16),
    ("file_size", "file_size", 11),
    ("format", "format", 8),
    ("dimensions", "dimensions", 12),
    ("needs_review", "needs_review", 13),
]

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F4E4A")
HDR_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
CELL_FONT = Font(name=ARIAL, size=10)
CHK_FILL = PatternFill("solid", fgColor="FDF6E3")   # confirm the match
LBL_FILL = PatternFill("solid", fgColor="E8F1E9")   # name read from the label
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()

# ------------------------------------------------------------------- Catalog
ws = wb.active
ws.title = "Catalog"
for c, (hdr, _, width) in enumerate(COLS, 1):
    cell = ws.cell(row=1, column=c, value=hdr)
    cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
    ws.column_dimensions[get_column_letter(c)].width = width
ws.row_dimensions[1].height = 22

for r, row in enumerate(rows, 2):
    for c, (_, key, _) in enumerate(COLS, 1):
        v = row.get(key, "")
        if key == "file_size" and v:
            v = int(v)
        cell = ws.cell(row=r, column=c, value=v)
        cell.font, cell.border = CELL_FONT, BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(key in ("description", "label_text")))
    fill = None
    if row["needs_review"] == "check":
        fill = CHK_FILL
    elif row["evidence_source"].startswith("label"):
        fill = LBL_FILL
    if fill:
        for c in range(1, len(COLS) + 1):
            ws.cell(row=r, column=c).fill = fill

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(rows) + 1}"

# --------------------------------------------------- Not On Website (label-only)
ws2 = wb.create_sheet("Not On Website")
lab = [r for r in rows if r["site_listed"] == "FALSE"]
prods = collections.OrderedDict()
for r in lab:
    prods.setdefault(r["product_name"], {"n": 0, "label": r["label_text"]})
    prods[r["product_name"]]["n"] += 1
note = ws2.cell(row=1, column=1, value=(
    "Real Kairali products that are NOT listed on kairaliproducts.in or .com. "
    "The name was read from the product's own packaging in the photo (Drive thumbnail + "
    "Drive's OCR of the label). Nothing here was guessed from a filename."))
note.font = Font(name=ARIAL, italic=True, size=10)
ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
ws2.row_dimensions[1].height = 30
for c, (hdr, width) in enumerate([("product_name", 40), ("images", 9), ("label_text (read from pack)", 95)], 1):
    cell = ws2.cell(row=2, column=c, value=hdr)
    cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
    ws2.column_dimensions[get_column_letter(c)].width = width
for i, (p, d) in enumerate(sorted(prods.items()), 3):
    for c, v in enumerate([p, d["n"], d["label"]], 1):
        cell = ws2.cell(row=i, column=c, value=v)
        cell.font, cell.border, cell.fill = CELL_FONT, BORDER, LBL_FILL
ws2.freeze_panes = "A3"

# ------------------------------------------------------------- Confirm Match
ws3 = wb.create_sheet("Confirm Match")
sub = [r for r in rows if r["needs_review"] == "check"]
n3 = ws3.cell(row=1, column=1, value=(
    "Probable site matches worth a human glance — the product IS on the website, but the "
    "filename wasn't an exact match. Nothing here is unidentified."))
n3.font = Font(name=ARIAL, italic=True, size=10)
ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
sc = [("new_filename", "proposed_new_filename", 42), ("original_drive_filename", "original_drive_filename", 36),
      ("product_name", "product_name", 32), ("matched_source_url", "matched_source_url", 60)]
for c, (hdr, _, width) in enumerate(sc, 1):
    cell = ws3.cell(row=2, column=c, value=hdr)
    cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
    ws3.column_dimensions[get_column_letter(c)].width = width
for i, row in enumerate(sorted(sub, key=lambda x: x["product_name"]), 3):
    for c, (_, key, _) in enumerate(sc, 1):
        cell = ws3.cell(row=i, column=c, value=row[key])
        cell.font, cell.border, cell.fill = CELL_FONT, BORDER, CHK_FILL
ws3.freeze_panes = "A3"

# ------------------------------------------------------------------- Summary
ws0 = wb.create_sheet("Summary", 0)
ws0.column_dimensions["A"].width = 44
ws0.column_dimensions["B"].width = 12
ws0.column_dimensions["C"].width = 74
t = ws0.cell(row=1, column=1, value="Kairali product images - Drive to project mapping")
t.font = Font(name=ARIAL, bold=True, size=14)
for i, txt in enumerate([
    "Source: Google Drive folder 'Pictures' (526) + subfolder 'group' (57) = 583 images.",
    "Names verified against kairaliproducts.in (primary, 137 products) then .com (fallback, 157).",
    "Where a product isn't on either site, the name was READ FROM THE PACKAGING in the photo.",
], 2):
    ws0.cell(row=i, column=1, value=txt).font = Font(name=ARIAL, size=10, italic=True)

last = len(rows) + 1
stats = [
    ("Total images mapped", f'=COUNTA(Catalog!A2:A{last})', "526 in Pictures + 57 in Pictures/group"),
    ("", "", ""),
    ("Named from the website", f'=COUNTIF(Catalog!H2:H{last},"site:kairaliproducts.in")', "Matched to a live product page"),
    ("Named from website + label", f'=COUNTIF(Catalog!H2:H{last},"label + site:kairaliproducts.in")', "Filename was wrong; the label identified it"),
    ("Named from the label only", f'=COUNTIF(Catalog!H2:H{last},"label (read from packaging)")', "Real product, NOT sold on either site"),
    ("Group shots / mock-ups", f'=COUNTIF(Catalog!H2:H{last},"n/a (not a single-product photo)")', "Multiple products per frame; no single name"),
    ("STILL UNIDENTIFIED", f'=COUNTIF(Catalog!H2:H{last},"none - unidentified")', "Nothing left unknown"),
    ("", "", ""),
    ("Listed on the website", f'=COUNTIF(Catalog!G2:G{last},"TRUE")', "site_listed = TRUE"),
    ("NOT listed on the website", f'=COUNTIF(Catalog!G2:G{last},"FALSE")', "see the 'Not On Website' tab"),
    ("Matches worth confirming", f'=COUNTIF(Catalog!R2:R{last},"check")', "see the 'Confirm Match' tab"),
]
for c, h in enumerate(["Metric", "Count", "Notes"], 1):
    cell = ws0.cell(row=6, column=c, value=h)
    cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
for i, (label, formula, note_) in enumerate(stats, 7):
    a = ws0.cell(row=i, column=1, value=label); a.font = Font(name=ARIAL, size=10, bold=bool(label))
    b = ws0.cell(row=i, column=2, value=formula); b.font = CELL_FONT
    c_ = ws0.cell(row=i, column=3, value=note_); c_.font = Font(name=ARIAL, size=10, italic=True)
    for c in range(1, 4):
        ws0.cell(row=i, column=c).border = BORDER

ws0.cell(row=19, column=1, value="file_size and dimensions are filled in by fetch_and_organize.py when the "
                                 "images are downloaded (they can only be read from the bytes).").font = \
    Font(name=ARIAL, size=10, italic=True, color="808080")

out = os.path.join(D, "kairali_image_catalog.xlsx")
wb.save(out)
print("wrote", out, "| label-only products:", len(prods))
